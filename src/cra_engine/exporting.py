"""Immutable, provenance-rich research packages; no external publication."""
import csv
import html
import io
import json
import re
import shutil
import sys
import uuid
import zipfile
from pathlib import Path

from .common import atomic, digest, dumps, now, require, save, sha, within


def export_diagnostic(store):
    from .runtime import status, sync
    error = None
    try:
        sync(store)
    except Exception as exc:
        from .common import EngineError
        error = str(exc) if isinstance(exc, EngineError) else "Input material could not be read; inspect locally"
    state = status(store)
    outcome = "infeasibility_diagnostic" if state["outcome"] == "infeasibility_diagnostic" else "waiting_external_condition"
    report = {"schema_version": 2, "research_completed": False, "outcome": outcome, "input_problem": error,
              "state": state, "recovery_conditions": "Resolve recorded input/method/permission/budget issues, then revalidate and obtain current human approvals", "synthetic": store.meta("synthetic")}
    target = store.root / "exports" / ("diagnostic-" + uuid.uuid4().hex[:12])
    target.mkdir(parents=True)
    save(target / "diagnostic.json", report)
    save(target / "specification-ledger.json", [json.loads(r[0]) for r in store.db.execute("SELECT body FROM objects WHERE kind='specification_attempt' ORDER BY rowid")])
    save(target / "runs.json", [json.loads(r[0]) for r in store.db.execute("SELECT body FROM objects WHERE kind='run_record' ORDER BY rowid")])
    atomic(target / "README.md", "# Research diagnostic handoff\n\nResearch is incomplete. No manuscript or scientific finding is certified.\n\nOutcome: " + outcome + "\n\nSee diagnostic.json for unresolved evidence, state, budget and recovery conditions. Raw data and credentials are not included.\n")
    save(target / "manifest.json", {"schema_version": 2, "kind": "diagnostic", "working": True, "input_included": False,
                                   "files": {p.name: sha(p) for p in target.iterdir() if p.is_file()}})
    archive = target.with_suffix(".zip")
    with zipfile.ZipFile(archive, "x", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(target.iterdir()):
            bundle.write(path, path.name)
    with store.tx():
        store.event("diagnostic_exported", {"path": archive.relative_to(store.root).as_posix(), "outcome": outcome})
    return {"directory": str(target), "archive": str(archive), "kind": "diagnostic", "research_completed": False}


def exhibits(analysis):
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(["model", "family", "n", "term", "estimate", "se", "ci_low", "ci_high", "p", "covariance"])
    entries = []
    for model in analysis["models"]:
        for p in model["parameters"]:
            writer.writerow([model["id"], model["family"], model["n"], p["term"], *[format(p[k], ".12g") for k in ("estimate", "se", "ci_low", "ci_high", "p")], model["covariance"]])
            entries.append((model["id"] + ":" + p["term"], p))
    lows = [p["ci_low"] for _, p in entries] + [0]
    highs = [p["ci_high"] for _, p in entries] + [0]
    low, high = min(lows), max(highs)
    if high == low:
        high += 1
    def x(v):
        return 240 + 480 * (v - low) / (high - low)
    height = 80 + 34 * len(entries)
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="780" height="%d" viewBox="0 0 780 %d">' % (height, height),
           '<rect width="100%" height="100%" fill="white"/>', '<desc>Coefficients and 95 percent intervals, generated directly from analysis.json; associations only.</desc>',
           '<line x1="%.3f" x2="%.3f" y1="20" y2="%d" stroke="gray" stroke-dasharray="4 4"/>' % (x(0), x(0), height - 35)]
    for index, (label, p) in enumerate(entries):
        y = 35 + 34 * index
        svg.append('<text x="5" y="%d" font-size="12">%s</text>' % (y + 4, html.escape(label)))
        svg.append('<line x1="%.6f" x2="%.6f" y1="%d" y2="%d" stroke="black"/>' % (x(p["ci_low"]), x(p["ci_high"]), y, y))
        svg.append('<circle cx="%.6f" cy="%d" r="3" data-estimate="%.12g" data-low="%.12g" data-high="%.12g"/>' % (x(p["estimate"]), y, p["estimate"], p["ci_low"], p["ci_high"]))
    svg += ['<text x="240" y="%d" font-size="12">%.4g</text>' % (height - 12, low), '<text x="685" y="%d" font-size="12">%.4g</text>' % (height - 12, high), '</svg>']
    return output.getvalue(), "\n".join(svg)


def make_docx(text, analysis, destination, language="zh-CN"):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    cjk_font = "Hiragino Sans GB" if sys.platform == "darwin" else "Noto Serif CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk_font)
        ppr = style._element.get_or_add_pPr()
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
    for line in text.splitlines():
        if not line.strip():
            continue
        line = re.sub(r"\[\[(?:stat|ref):[^\]]+\]\]", "", line)
        line = re.sub(r"\[\[cite:([^\]]+)\]\]", r"[\1]", line)
        line = line.replace("[[Table:T1]]", "Table 1" if language.startswith("en") else "表1").replace("[[Figure:F1]]", "Figure 1" if language.startswith("en") else "图1")
        if line.startswith("# "):
            paragraph = document.add_paragraph(line[2:], "Title")
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        else:
            document.add_paragraph(line.replace("**", ""))
    document.add_heading("Table 1 Model results" if language.startswith("en") else "表1 模型结果", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Model/term" if language.startswith("en") else "模型/变量", "N", "b", "SE", "95% CI", "p"]):
        cell.text = label
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for model in analysis["models"]:
        for p in model["parameters"]:
            row = table.add_row().cells
            values = [model["id"] + "/" + p["term"], str(model["n"]), "%.4f" % p["estimate"], "%.4f" % p["se"], "[%.4f, %.4f]" % (p["ci_low"], p["ci_high"]), "%.4g" % p["p"]]
            for cell, value in zip(row, values):
                cell.text = value
    document.add_paragraph("Figure 1 and its source values are in figure.svg and analysis.json. All approved specifications are retained, irrespective of statistical significance." if language.startswith("en") else "图1及其数值源另见 figure.svg 和 analysis.json。表中保留全部已执行并批准的规格，不以显著性筛选。")
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.save(destination)


def export_package(store, working=False, include_data=False):
    from .runtime import readiness, sync
    sync(store)
    ready = readiness(store)
    require(store.object_hash("draft") and store.object_hash("analysis"), "No fresh manuscript and analysis to export")
    require(working or ready["ready_for_final"] and store.approved("final"), "Final export requires resolved blockers and current human final approval; use --working for a clearly marked draft")
    ident = digest({"draft": store.object_hash("draft"), "analysis": store.object_hash("analysis"), "at": now()})[:16]
    target = store.root / "exports" / ident
    target.mkdir(parents=True)
    analysis = store.get("analysis", fresh=True)
    draft = store.get("draft", fresh=True)
    text = draft["text"]
    language = draft.get("language", "zh-CN")
    if working:
        text = ("# Unapproved working draft\n\n" if language.startswith("en") else "# 未终审工作稿\n\n") + "This working draft is not approved for submission.\n\n" + text
    atomic(target / "manuscript.md", text)
    make_docx(text, analysis, target / "manuscript.docx", language)
    table, figure = exhibits(analysis)
    atomic(target / "table.csv", table)
    atomic(target / "figure.svg", figure)
    for ident in ("analysis", "protocol", "claims", "corpus", "method_review", "evidence_review", "audit"):
        if store.object_hash(ident):
            save(target / (ident + ".json"), store.get(ident, fresh=True))
    specs = [json.loads(r[0]) for r in store.db.execute("SELECT body FROM objects WHERE kind='specification_attempt' ORDER BY rowid")]
    save(target / "specification-ledger.json", specs)
    save(target / "runs.json", [json.loads(r[0]) for r in store.db.execute("SELECT body FROM objects WHERE kind='run_record' ORDER BY rowid")])
    save(target / "model-calls.json", [dict(r) for r in store.db.execute("SELECT id,task,status,reserved,actual,model,usage,at FROM calls ORDER BY rowid")])
    save(target / "limitations.json", {"working": working, "readiness": ready, "privacy": "Data omitted unless explicitly included; aggregation is not an anonymity guarantee", "format": "DOCX uses internal reference keys; complete APA7/target formatting and anonymous metadata require human verification", "real_project_pilots": "not certified by software tests"})
    save(target / "budget.json", store.budget())
    brief = store.get("brief", fresh=True)
    # Preserve source bytes: reserializing a canonical DB object changes its file hash.
    shutil.copyfile(within(store.root, brief["data"]["codebook"]), target / "codebook.json")
    if include_data:
        shutil.copyfile(within(store.root, brief["data"]["csv"]), target / "input.csv")
    # Ship the actual recipe sources and environment; no credentials or control DB.
    recipe_dir = target / "recipe/cra_engine"
    recipe_dir.mkdir(parents=True)
    for name in ("__init__.py", "common.py", "statistics.py", "worker.py", "isolation.py"):
        shutil.copyfile(Path(__file__).parent / name, recipe_dir / name)
    atomic(target / "reproduce.py", '''import sys,json,hashlib,math
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent/"recipe"))
from cra_engine.statistics import fit_models
p=Path(__file__).parent
data=Path(sys.argv[1]) if len(sys.argv)>1 else p/"input.csv"
expected=json.loads((p/"analysis.json").read_text())
if hashlib.sha256(data.read_bytes()).hexdigest()!=expected["data_sha256"]:
    raise SystemExit("Input hash differs; do not substitute untracked data")
plan=json.loads((p/"protocol.json").read_text())
result=fit_models(data,p/"codebook.json",plan,min_cell=expected["privacy_min_cell"])
(p/"reproduced.json").write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False))
def same(a,b):
    if type(a) in (int,float) and type(b) in (int,float): return math.isclose(a,b,rel_tol=1e-7,abs_tol=1e-9)
    if isinstance(a,dict) and isinstance(b,dict): return set(a)==set(b) and all(same(a[k],b[k]) for k in a)
    if isinstance(a,list) and isinstance(b,list): return len(a)==len(b) and all(same(x,y) for x,y in zip(a,b))
    return a==b
if not same(expected,result): raise SystemExit("Reproduction differs; inspect reproduced.json and pinned environment")
print("Reproduction PASS; all registered statistics and environment matched")
''')
    atomic(target / "requirements.txt", "\n".join(k + "==" + v for k, v in analysis["environment"].items()) + "\n")
    atomic(target / "README.md", "# Research package\n\n" + ("Working draft; final human approval pending.\n\n" if working else "Human approval recorded locally; not identity or scientific certification.\n\n") + "Use an isolated local environment. Install requirements.txt explicitly. Python version: " + analysis.get("recipe_identity", {}).get("python_version", "see analysis environment") + ". Run python reproduce.py /authorized/input.csv. Input SHA256: " + analysis["data_sha256"] + ". Numeric tolerance: relative 1e-7, absolute 1e-9. Model prose is not guaranteed byte reproducible. Internal citation keys require target-format editing.\n")
    files = {p.relative_to(target).as_posix(): sha(p) for p in sorted(target.rglob("*")) if p.is_file()}
    save(target / "manifest.json", {"schema_version": 2, "files": files, "input_included": include_data, "data_sha256": analysis["data_sha256"], "working": working})
    archive = target.with_suffix(".zip")
    with zipfile.ZipFile(archive, "x", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(target.rglob("*")):
            if path.is_file():
                bundle.write(path, path.relative_to(target))
    with store.tx():
        store.event("exported", {"path": archive.relative_to(store.root).as_posix(), "working": working, "include_data": include_data})
    return {"directory": str(target), "archive": str(archive), "working": working}
